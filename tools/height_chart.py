#!/usr/bin/env python3
"""캐릭터 키 비교 그림 (docs/screenshots/roster_heights.png): 게임과 같은 규칙 — 화면 크기 = 키(cm) / character.json figure_px.
모자 쓴 캐릭터는 figure_px 가 '모자 절반 높이 ~ 발' (tools/char_sprites.py hat). 키는 data/outgame.xlsx 대원 시트 키(cm)."""
import json, os
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
K = 3.2                       # 1cm = 3.2px
GROUND, LEFT, COLW = 720, 90, 250


def main():
    ws = load_workbook(os.path.join(ROOT, "data", "outgame.xlsx"), read_only=True)["대원"]
    rows = list(ws.iter_rows(values_only=True))
    hd = rows[0]
    chars = []
    for r in rows[1:]:
        d = dict(zip(hd, r))
        if d.get("스프라이트"):
            chars.append((str(d["이름"]), str(d["스프라이트"]), float(d["키(cm)"])))
    chars.sort(key=lambda c: -c[2])
    W, H = LEFT + COLW * len(chars) + 20, GROUND + 90
    o = Image.new("RGB", (W, H), (36, 38, 46))
    d = ImageDraw.Draw(o)
    f = ImageFont.truetype(os.path.join(ROOT, "assets/fonts/Pretendard-Bold.ttf"), 26)
    fs = ImageFont.truetype(os.path.join(ROOT, "assets/fonts/Pretendard-Medium.ttf"), 20)
    for cm in range(0, 200, 10):
        y = GROUND - cm * K
        major = cm % 50 == 0
        d.line([(LEFT - 10, y), (W - 10, y)], fill=(98, 104, 120) if major else (64, 68, 80), width=1)
        if major or cm >= 160:
            d.text((10, y - 12), "%dcm" % cm, fill=(150, 156, 170), font=fs)
    for i, (name, cid, cm) in enumerate(chars):
        info = json.load(open(os.path.join(ROOT, "assets/characters", cid, "character.json"), encoding="utf-8"))
        b = Image.open(os.path.join(ROOT, "assets/characters", cid, "base.png")).convert("RGBA")
        s = cm * K / float(info["figure_px"])           # 게임과 같은 배율
        im = b.resize((round(b.width * s), round(b.height * s)), Image.LANCZOS)
        gx, gy = info["ground"]
        cx = LEFT + i * COLW + COLW // 2
        o.paste(im, (round(cx - gx * s), round(GROUND - gy * s)), im)
        top = GROUND - cm * K
        d.line([(cx - 70, top), (cx + 70, top)], fill=(255, 210, 120), width=2)
        label = "%s  %s" % (name, ("%g" % cm))
        if "hat" in info:
            label += " (모자 절반)"
        tw = d.textlength(label, font=f if "hat" not in info else fs)
        d.text((cx - tw / 2, GROUND + 18), label, fill=(255, 230, 160), font=f if "hat" not in info else fs)
    d.line([(LEFT - 10, GROUND), (W - 10, GROUND)], fill=(200, 200, 210), width=3)
    out = os.path.join(ROOT, "docs", "screenshots", "roster_heights.png")
    o.save(out)
    print("saved", out, o.size)


if __name__ == "__main__":
    main()
