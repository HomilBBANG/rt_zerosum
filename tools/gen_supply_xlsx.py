#!/usr/bin/env python3
"""보급 · 납품 · 장군 데이터 엑셀(data/supply.xlsx) 초기 생성. 처음 한 번만 — 이후에는 엑셀에서 고친 파일이 원본.
주의: 다시 실행하면 고친 값을 덮어씀. (사용: python3 tools/gen_supply_xlsx.py --force)
모든 값은 [임시] (아이템 이름 · 확률 · 납품 목록 · 금액)."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "supply.xlsx")
if os.path.exists(OUT) and "--force" not in sys.argv:
    sys.exit("data/supply.xlsx 가 이미 있습니다 (고친 값이 사라지므로 덮어쓰려면 --force)")

SETTINGS = [
    ["id", "값", "설명"],
    ["base_pay", 150, "기본 주급 USD — 결산 때 주급 = 기본 주급 × 그 주 납품 비율 (50% 납품 → 75) [임시]"],
    ["general_bonus", 40, "그 주 장군 요구에 물건을 주면 결산 때 추가금 USD [임시]"],
    ["general_penalty", 40, "그 주 장군 요구를 거절(또는 결산까지 대답 안 함)하면 결산 때 차감 USD [임시]"],
    ["loyalty_start", 50, "장군 충성심 시작 값 (0~100) [임시]"],
    ["loyalty_give", 5, "장군에게 물건을 주면 충성심 + [확정]"],
    ["loyalty_refuse", -10, "장군 요구를 거절하면 충성심 (음수) [확정]"],
]

ITEMS = [
    ["id", "이름", "설명"],
    ["scrap", "고철", "부서진 장비에서 나온 쓸 만한 쇳조각 [임시]"],
    ["gunpowder", "화약", "불발탄에서 긁어모은 화약 [임시]"],
    ["ammo", "탄약 상자", "아직 뜯지 않은 탄약 한 상자 [임시]"],
    ["bandage", "붕대", "깨끗한 붕대 한 묶음 [임시]"],
    ["ration", "전투 식량", "밀봉된 전투 식량 [임시]"],
    ["ghoul_claw", "구울 발톱", "구울에게서 떼어 낸 발톱 — 연구용으로 수거 [임시]"],
    ["wolf_pelt", "늑대 가죽", "변이 늑대의 질긴 가죽 [임시]"],
    ["slime_gel", "슬라임 점액", "병에 담은 슬라임 점액 [임시]"],
    ["old_coin", "옛 동전", "무너진 도시에서 주운 옛 동전 [임시]"],
    ["cigarette", "담배", "귀한 담배 한 갑 [임시]"],
    ["whiskey", "위스키", "전쟁 전 병 위스키 [임시]"],
    ["map_piece", "지도 조각", "옛 군용 지도의 한 조각 [임시]"],
]

BATTLE = [
    ["스테이지", "아이템", "최소", "최대", "확률%", "설명"],
    ["*", "scrap", 1, 3, 100, "* = 모든 전투 승리 [임시]"],
    ["*", "gunpowder", 1, 2, 60, ""],
    ["*", "ammo", 1, 1, 30, ""],
    ["*", "old_coin", 1, 1, 20, ""],
    ["stage_01", "ghoul_claw", 1, 2, 70, "그 스테이지에서 이기면 [임시]"],
    ["stage_01", "wolf_pelt", 1, 1, 50, ""],
    ["stage_01", "slime_gel", 1, 2, 50, ""],
]

GIFTS = [
    ["대원", "아이템", "확률%", "최소친밀도", "설명"],
    ["npc_dan", "cigarette", 25, 0, "하루 첫 대화 때 확률로 건넴 [임시]"],
    ["npc_kira", "ammo", 20, 0, ""],
    ["npc_ethan", "map_piece", 15, 0, ""],
    ["npc_hana", "ration", 25, 0, ""],
    ["npc_alma", "bandage", 30, 0, ""],
    ["npc_nikita", "whiskey", 20, 30, "친밀도 30 이상일 때만 [임시]"],
    ["npc_rohan", "gunpowder", 25, 0, ""],
]

DEMANDS = [
    ["아이템", "개수", "가중치", "장군 대사"],
    ["whiskey", 1, 1, "자네 기지에 좋은 술이 있다던데. 한 병 내놓게."],
    ["cigarette", 2, 1, "담배가 떨어졌네. 두 갑만 챙겨 주게."],
    ["wolf_pelt", 1, 1, "겨울이 오기 전에 늑대 가죽 하나 구해 두게."],
    ["old_coin", 1, 1, "옛 동전을 모으는 게 내 취미라네. 하나 있나?"],
    ["ghoul_claw", 2, 1, "본부에서 구울 표본을 원하네. 발톱 두 개."],
    ["ration", 3, 1, "내 부하들 먹일 전투 식량 세 개."],
]

DELIVERY = [
    ["진행률%", "납품 목록", "설명"],
    [0, "scrap:4, gunpowder:2, bandage:1", "스토리 진행률(완료한 퀘스트 비율)이 이 값 이상일 때. 같은 진행률 줄이 여럿이면 주마다 번갈아 [임시]"],
    [0, "scrap:3, ammo:1, ration:2", ""],
    [50, "scrap:5, ghoul_claw:2, gunpowder:3", ""],
    [50, "scrap:4, wolf_pelt:1, slime_gel:2, ammo:1", ""],
    [100, "scrap:6, ghoul_claw:3, wolf_pelt:2, gunpowder:3", ""],
]

README = [
    ["보급 · 납품 · 장군 (data/supply.xlsx)", ""],
    ["", ""],
    ["설정", "기본 주급 · 장군 추가금/차감 · 충성심"],
    ["아이템", "전리품 — 전투 승리 보상이나 대원에게 얻음. 납품 · 장군 요구에 씀"],
    ["전투 보상", "스테이지마다(* = 모든 승리) 아이템 · 개수 · 확률"],
    ["대원 선물", "그 대원과 하루 첫 대화 때 확률로 아이템을 줌 (최소 친밀도)"],
    ["장군 요구", "장군 점검 때 요구하는 아이템 후보 (지난번과 다른 것을 가중치로 고름) · 대사"],
    ["납품", "주마다 납품 목록 — 스토리 진행률(완료한 퀘스트 비율 %) 이상인 가장 높은 줄. 같은 진행률 줄이 여럿이면 주마다 번갈아"],
    ["", ""],
    ["주급", "결산 때 기본 주급 × 납품 비율 (예: 150 × 50% = 75). 납품 비율 = 목록 개수 합 중 납품한 개수"],
    ["장군", "매달 한 번 점검 → 아이템 요구 → 준다(결산 추가금 · 충성심 +5) / 안 준다(결산 차감 · 충성심 -10). 결산까지 대답 안 하면 거절"],
    ["실시간", "게임 실행 중 저장하면 다시 읽음 (기지 데이터와 같이)"],
]


def sheet(wb, name, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    for r in rows:
        ws.append(r)
    head = PatternFill("solid", fgColor="3B4A2F")
    if name != "읽어보기":
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
        ws.freeze_panes = "A2"
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


wb = Workbook()
sheet(wb, "설정", SETTINGS, [18, 10, 90], True)
sheet(wb, "아이템", ITEMS, [14, 14, 60])
sheet(wb, "전투 보상", BATTLE, [12, 14, 8, 8, 8, 40])
sheet(wb, "대원 선물", GIFTS, [12, 14, 8, 12, 40])
sheet(wb, "장군 요구", DEMANDS, [14, 8, 8, 60])
sheet(wb, "납품", DELIVERY, [10, 50, 70])
sheet(wb, "읽어보기", README, [26, 100])
wb.save(OUT)
print("→", OUT)
