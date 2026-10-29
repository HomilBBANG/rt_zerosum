#!/usr/bin/env python3
"""밸런스 엑셀(data/balance.xlsx) 초기 생성. 처음 한 번만 — 이후에는 엑셀에서 직접 고치는 파일이 원본.
주의: 다시 실행하면 엑셀에서 고친 값을 덮어씀. (사용: python3 tools/gen_balance_xlsx.py --force)"""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "balance.xlsx")
if os.path.exists(OUT) and "--force" not in sys.argv:
    sys.exit("data/balance.xlsx 가 이미 있습니다 (엑셀에서 고친 값이 사라지므로 덮어쓰려면 --force)")

CHAR_COLS = ["id", "이름", "병과", "레벨", "속도", "사거리", "HP", "AP", "AP회복", "방어", "스킬", "무기", "열", "줄", "키(cm)", "스프라이트", "눈색", "원화방향"]
CHARS = [
    ["a_dan", "단", "돌격병", 5, 8, "근접", 46, 6, 2, 2, "basic_attack, heavy_attack, ultimate", "axe, pistol", "전방", 1, 185, "dan", "#7fb2ff", "오른쪽"],
    ["a_kira", "키라", "돌격병", 5, 10, "근접", 40, 6, 2, 1, "basic_attack, heavy_attack, ultimate", "combat_knife, pistol", "전방", 2, 175, "kira", "#c8d4ff", "오른쪽"],
    ["a_ethan", "이선", "소총수", 5, 9, "원거리", 30, 6, 2, 0, "basic_attack, heavy_attack, ultimate", "rifle, bayonet", "중앙", 1, 178, "ethan", "", "오른쪽"],
    ["a_hana", "하나", "돌격소총수", 5, 9, "원거리", 34, 6, 2, 1, "basic_attack, heavy_attack, ultimate", "rifle, bayonet", "중앙", 2, 176.8, "hero2", "#ffb020", "오른쪽"],
    ["a_alma", "앨마", "의무병", 4, 7, "원거리", 28, 7, 3, 0, "basic_attack, heal, rally, ultimate", "pistol, syringe", "후방", 1, 166, "alma", "#43d6a0", "오른쪽"],
    ["a_nikita", "니키타", "저격수", 6, 11, "원거리", 24, 5, 2, 0, "basic_attack, heavy_attack, ultimate", "sniper, knife", "후방", 2, 168, "nikita_army", "#4f9cff", "오른쪽"],
]
MON_COLS = ["id", "이름", "분류", "레벨", "속도", "사거리", "HP", "AP", "AP회복", "방어", "무기", "차지칸", "키(cm)", "스프라이트", "원화방향", "AI"]
MONS = [
    ["ghoul", "구울", "괴물", 6, 10, "근접", 110, 6, 2, 2, "claw", 2, 240, "ghoul", "왼쪽", "default"],
    ["wolf", "늑대", "괴물", 7, 12, "근접", 90, 6, 2, 1, "claw", 1, 400, "wolf", "오른쪽", "aggressive"],
    ["slime", "슬라임", "괴물", 4, 5, "근접", 130, 5, 2, 3, "slam", 2, 400, "slime", "오른쪽", "default"],
]
ITEM_COLS = ["id", "이름", "종류", "사거리", "위력보너스", "스프라이트", "설명"]
ITEMS = [
    ["combat_knife", "전투단검", "무기", "근접", 2, "knife", ""],
    ["knife", "단검", "무기", "근접", 0, "knife", ""],
    ["pistol", "권총", "무기", "원거리", 0, "pistol", ""],
    ["sniper", "저격총", "무기", "원거리", 3, "sniper", ""],
    ["rifle", "소총", "무기", "원거리", 1, "rifle", ""],
    ["bayonet", "대검", "무기", "근접", 2, "bayonet", ""],
    ["axe", "전투도끼", "무기", "근접", 2, "axe", ""],
    ["syringe", "전투 주사기", "무기", "근접", 1, "syringe", ""],
    ["claw", "발톱", "무기", "근접", 3, "", "괴물 전용"],
    ["slam", "덮치기", "무기", "근접", 2, "", "괴물 전용"],
    ["shotgun", "산탄총", "무기", "원거리", 2, "shotgun", "연구로 해금 (더미 수치)"],
    ["smg", "기관단총", "무기", "원거리", 1, "smg", "연구로 해금 (더미 수치)"],
    ["blade", "고주파 블레이드", "무기", "근접", 4, "blade", "연구로 해금 (더미 수치)"],
    ["mg", "경기관총", "무기", "원거리", 3, "mg", "연구로 해금 (더미 수치)"],
    ["launcher", "유탄발사기", "무기", "원거리", 4, "launcher", "연구로 해금 (더미 수치)"],
]
SKILL_COLS = ["id", "이름", "종류", "AP비용", "대상", "효과", "위력", "설명"]
SKILLS = [
    ["basic_attack", "기본 공격", "attack", 1, "weapon", "damage", 6, "현재 무기로 공격. 근접: 상대 앞열 우선 / 원거리: 전열 무관"],
    ["heavy_attack", "강공격", "heavy", 2, "weapon", "damage", 11, "현재 무기로 강하게 공격"],
    ["heal", "응급 처치", "support", 2, "ally_any", "heal", 10, "아군 1명 HP 회복"],
    ["rally", "독려", "support", 2, "ally_other", "ap_grant", 2, "다른 아군 1명 AP +2 (상한까지)"],
    ["ultimate", "필살기", "ultimate", 4, "enemy_any", "damage", 18, "열 무관 강력한 일격"],
    ["move", "위치 변경", "move", 1, "slot", "move", 0, "빈 슬롯 이동 / 점유 슬롯 아군과 교환"],
    ["weapon", "무기 변경", "weapon", 1, "none", "weapon_swap", 0, "보조 무기로 교체 (사거리 변경)"],
    ["brutal_strike", "강타", "heavy", 3, "weapon", "damage", 16, "스킬 트리(액티브): 현재 무기로 아주 강하게 공격"],
    ["flank", "측면 기습", "attack", 2, "enemy_any", "damage", 9, "스킬 트리(액티브): 열 무관 적 1명 공격"],
    ["second_wind", "재정비", "support", 1, "ally_any", "heal", 8, "스킬 트리(액티브): 아군 1명(자신 포함) HP 회복"],
    ["inspire", "고무", "support", 1, "ally_other", "ap_grant", 1, "스킬 트리(액티브): 다른 아군 1명 AP +1"],
    ["incendiary", "소이탄", "heavy", 2, "weapon", "damage", 6, "스킬 트리(액티브): 맞으면 화상 (화염 속성 고정 피해)"],
    ["suppress", "제압 사격", "attack", 2, "enemy_any", "damage", 5, "스킬 트리(액티브): 열 무관, 맞으면 감속"],
    ["painkiller_shot", "진통제 투여", "support", 2, "ally_any", "status", 0, "[임시] 아군 1명에게 진통제 — 차례마다 최대 HP 회복 (상태 '진통제' = 효과 재생)"],
    ["self_heal", "자가 치유", "support", 1, "self", "status", 0, "[임시] 자신 — 차례마다 최대 HP 회복 (상태 '자가 치유' = 효과 재생)"],
]
# [임시] 명중 · 회피 (시트 맨 끝 열 — 열 순서는 자유, 머리글로 찾음)
_HIT = {"a_dan": (0, 5), "a_kira": (0, 15), "a_ethan": (10, 5), "a_hana": (5, 8), "a_alma": (0, 8), "a_nikita": (15, 10),
        "ghoul": (0, 5), "wolf": (0, 20), "slime": (-5, 0)}
CHAR_COLS = CHAR_COLS + ["명중", "회피"]
MON_COLS = MON_COLS + ["명중", "회피"]
CHARS = [r + list(_HIT.get(r[0], (0, 0))) for r in CHARS]
MONS = [r + list(_HIT.get(r[0], (0, 0))) for r in MONS]
_SKILL_ACC = {"basic_attack": 95, "heavy_attack": 85, "ultimate": 100, "brutal_strike": 80, "flank": 90, "incendiary": 85, "suppress": 90}
SKILL_COLS = SKILL_COLS + ["명중"]
SKILLS = [r + [_SKILL_ACC.get(r[0], "")] for r in SKILLS]
STAGE_COLS = ["id", "이름", "조명", "선공", "적HP배율", "적공격배율", "시드", "설명"]
STAGES = [
    ["stage_01", "첫 조우", "g1", "동전", 1.0, 1.0, 12345, "구울(전방·중앙) + 늑대(후방). 배치는 스테이지 에디터"],
]
README = [
    ["시트", "열", "설명"],
    ["스킬트리", "", "스킬 트리 노드는 엑셀이 아니라 data/skill_tree.json — Godot 에디터 '스킬 트리' 탭에서 짬 (캐릭터별 액티브·패시브 십자형 트리)"],
    ["신앙", "", "신 종류·신앙 스킬·누가 어느 신을 믿는지(시작 신앙도)는 엑셀이 아니라 data/faith.json — Godot 에디터 '스킬 트리' 탭 → '신앙' 에서 짬"],
    ["공통", "id", "다른 시트·파일에서 참조하는 고유 이름 (영문 권장, 바꾸면 참조도 같이 바꿀 것)"],
    ["캐릭터", "사거리", "근접 / 원거리 — 행동 순서 동률 정렬용 병과 기본 사거리"],
    ["캐릭터", "스킬 / 무기", "쉼표로 구분한 스킬 id / 아이템 id. 무기 첫 번째 = 주무기, 두세 번째 = 보조무기"],
    ["캐릭터", "열 / 줄", "기본 배치: 열 = 전방·중앙·후방, 줄 = 1(뒤쪽 줄)·2(카메라 쪽 줄)"],
    ["캐릭터", "눈색", "눈동자를 일부러 칠할 색(#RRGGBB). 비우면 원화 눈 색 그대로"],
    ["캐릭터·몬스터", "원화방향", "원화가 바라보는 쪽 (오른쪽 / 왼쪽) — 화면에서 상대를 보게 자동 반전"],
    ["몬스터", "차지칸", "1 = 한 칸, 2 = 전방+중앙처럼 연속 두 열 (이동·교환 불가)"],
    ["몬스터", "AI", "data/ai/<AI>.json — Godot 에디터 'AI' 탭에서 만들고 고침"],
    ["몬스터", "스킬", "몬스터가 쓰는 스킬·몬스터 전용 스킬은 엑셀이 아니라 data/monster_skills.json — Godot 에디터 '스킬 트리' 탭 → '몬스터 스킬' 에서 짬 (공용 스킬 = 이 파일 스킬 시트도 넣을 수 있음)"],
    ["아이템", "스프라이트", "assets/sprites/w_<스프라이트>.png (비우면 손에 든 무기 그림 없음)"],
    ["스킬", "종류", "attack / heavy / support / ultimate / move / weapon (ultimate = 필살기 칸)"],
    ["스킬", "대상", "weapon(무기 사거리) / enemy_any / ally_any / ally_other / self / slot / none"],
    ["스킬", "효과", "damage(피해) / heal(회복) / ap_grant(AP 부여) / status(상태 이상만) / move / weapon_swap"],
    ["스킬", "상태 이상", "스킬에 다는 상태 이상(출혈·화상·동상·감속·가속·기절·구속·실명, 턴·수치·확률)은 엑셀이 아니라 data/skill_status.json — Godot 에디터 '스킬 트리' 탭 → '상태 이상' 에서 담"],
    ["스킬", "상태 이상 적용 시점", "처음 걸릴 때 바로 한 번 + 그 캐릭터 차례가 올 때마다 행동보다 먼저. 차례가 끝날 때 상태턴 1 감소 (동상은 차례 시작에 감소, 0 이면 터짐). 다시 걸리면 남은 턴·수치 중 큰 쪽"],
    ["스테이지", "조명 / 선공", "조명 g1~g4, 선공 동전 / 기습(아군 선공) / 매복(적 선공)"],
    ["스테이지", "적HP배율 / 적공격배율", "이 스테이지의 적 최대 HP·피해에 곱함 (1 = 그대로)"],
    ["스테이지", "배치", "적 배치는 data/stages/<id>.json — 메뉴 → 스테이지 에디터"],
    ["캐릭터·몬스터", "명중 / 회피", "[임시] 명중 = 명중 보너스(%p, 쓰는 스킬 명중에 더함) · 회피 = 적의 공격을 피할 확률(%). 비우면 0"],
    ["스킬", "명중", "[임시] 적에게 쓰는 공격·상태 스킬의 기본 명중(%). 명중률 = 스킬 명중 + 쓰는 쪽 명중 - 맞는 쪽 회피 (5~100%). 비우면 공격 90 · 상태 이상만 85, 회복·AP·이동은 늘 성공. 기절·구속된 대상은 못 피함"],
    ["반영", "", "게임 실행 중 이 파일을 저장하면 1초 안에 다시 읽어 바로 적용 (현재 전투 수치 포함). 편성·배치 변경은 다음 전투부터"],
]

HEAD_FILL = PatternFill("solid", fgColor="2E3440")
wb = Workbook()
first = True
for name, cols, rows, widths in [
    ("캐릭터", CHAR_COLS, CHARS, None), ("몬스터", MON_COLS, MONS, None), ("아이템", ITEM_COLS, ITEMS, None),
    ("스킬", SKILL_COLS, SKILLS, None), ("스테이지", STAGE_COLS, STAGES, None), ("읽어보기", README[0], README[1:], [12, 22, 90])]:
    ws = wb.active if first else wb.create_sheet()
    first = False
    ws.title = name
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2" if name != "읽어보기" else "A2"
    for i, col in enumerate(cols):
        letter = ws.cell(row=1, column=i + 1).column_letter
        w = widths[i] if widths else max(8, min(40, max(len(str(col)) * 2 + 2, *[len(str(r[i])) + 2 for r in rows if i < len(r)])))
        ws.column_dimensions[letter].width = w
    # 선택지 칸은 목록으로
    def pick(colname, opts):
        if colname in cols:
            letter = ws.cell(row=1, column=cols.index(colname) + 1).column_letter
            dv = DataValidation(type="list", formula1='"%s"' % ",".join(opts), allow_blank=True)
            dv.add(f"{letter}2:{letter}200")
            ws.add_data_validation(dv)
    pick("사거리", ["근접", "원거리"])
    pick("열", ["전방", "중앙", "후방"])
    pick("원화방향", ["오른쪽", "왼쪽"])
    pick("조명", ["g1", "g2", "g3", "g4"])
    pick("선공", ["동전", "기습", "매복"])
    if name == "스킬":
        pick("종류", ["attack", "heavy", "support", "ultimate", "move", "weapon"])
        pick("효과", ["damage", "heal", "ap_grant", "move", "weapon_swap"])
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT)
